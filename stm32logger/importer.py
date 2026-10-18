"""Daha önce kaydedilmiş dosyaları (yeni TXT/CSV/XLSX ya da eski sürüm .txt) açar."""
from __future__ import annotations

import csv
import itertools
import math
import os

from .i18n import STRINGS
from .parser import parse_line, strip_old_timestamp, NUM_RE
from .session import MAX_CHANNELS, Session, _dedupe, clean_note

MAX_FILE_BYTES = 4 * 1024 ** 3     # 4 GB üstü dosya açılmaz (bellek koruması)

_META_KEYS = ("col_index", "col_seq", "col_time_s", "col_pc_time", "col_elapsed", "col_note")


def _meta_names() -> dict[str, str]:
    out = {}
    for lang in STRINGS.values():
        for k in _META_KEYS:
            out[lang[k].strip().lower()] = k
    return out


def _to_float(v):
    if v is None or v == "":
        return float("nan")
    if isinstance(v, (int, float)):
        f = float(v)
    else:
        v = str(v).strip()
        f = float(v) if NUM_RE.match(v) else float("nan")
    return f if math.isfinite(f) else float("nan")


def _session_from_table(heads: list[str], rows, channel_prefix: str) -> Session | None:
    meta = _meta_names()
    roles = [meta.get(str(h).strip().lower()) for h in heads]
    if not any(roles):
        return None
    s = Session(mode="file", seq_mode="off", channel_prefix=channel_prefix)
    ch_idx = [i for i, r in enumerate(roles) if r is None][:MAX_CHANNELS]
    def unguard(h: str) -> str:   # dışa aktarımdaki formül korumasını ('=x) geri al
        return h[1:] if h[:1] == "'" and h[1:2] in ("=", "+", "-", "@") else h

    names = _dedupe([unguard(str(heads[i]).strip()) or f"{channel_prefix} {j + 1}" for j, i in enumerate(ch_idx)])
    seq_i = roles.index("col_seq") if "col_seq" in roles else None
    time_i = roles.index("col_time_s") if "col_time_s" in roles else None
    el_i = roles.index("col_elapsed") if "col_elapsed" in roles else None
    note_i = roles.index("col_note") if "col_note" in roles else None
    for name in names:
        s._ensure(name)
    s.has_seq = seq_i is not None
    times = []
    for r in rows:
        if not r or all(c in (None, "") for c in r):
            continue
        r = list(r) + [None] * (len(heads) - len(r))
        for j, i in enumerate(ch_idx):
            s.cols[j].append(_to_float(r[i]))
        if seq_i is not None:
            v = _to_float(r[seq_i])
            sv = int(v) if v == v else s.n
            s.seq.append(s._check_seq(sv))
        tv = _to_float(r[el_i]) if el_i is not None else float(s.n)
        s.t_host.append(tv if tv == tv else float(s.n))
        if time_i is not None and len(times) < 2:
            times.append(_to_float(r[time_i]))
        if note_i is not None and r[note_i] not in (None, ""):
            note = clean_note(r[note_i])
            if note[:1] == "'" and note[1:2] in ("=", "+", "-", "@"):
                note = note[1:]
            if note:
                s.notes.append((s.n, tv if tv == tv else float(s.n), note))
        s.n += 1
    if len(times) == 2 and times[1] > times[0]:
        s.sample_rate = round(1.0 / (times[1] - times[0]), 6)
    if el_i is not None:
        s.mode = "live"   # x ekseni PC zamanı
    return s


def load_file(path: str, channel_prefix: str = "Kanal") -> Session:
    if os.path.getsize(path) > MAX_FILE_BYTES:
        raise ValueError("too-big")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        s = _load_xlsx(path, channel_prefix)
    elif ext == ".bin":
        s = _load_bin(path, channel_prefix)
    else:
        s = _load_text(path, channel_prefix)
    s.port = os.path.basename(path)
    s.finish("file")
    s.imported = True
    s.mark_saved()          # zaten bir dosyadan geldi: yeni not eklenmedikçe "kaydedilmedi" sayılmaz
    return s


def _load_xlsx(path: str, channel_prefix: str) -> Session:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            try:
                heads = [str(h) if h is not None else "" for h in next(it)]
            except StopIteration:
                continue
            sheets.append((heads, it))
        for k, (heads, it) in enumerate(sheets):
            # >1M satırlık kayıtlar birden çok "Veriler" sayfasına bölünür: aynı başlıklıları birleştir
            rest = [it2 for h2, it2 in sheets[k + 1:] if h2 == heads]
            s = _session_from_table(heads, itertools.chain(it, *rest), channel_prefix)
            if s is not None:
                return s
    finally:
        wb.close()
    raise ValueError("no-data")


def _lines(path: str):
    """Dosyayı satır satır okur (tamamını belleğe almaz)."""
    with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
        for ln in f:
            yield ln.rstrip("\r\n")


def _load_bin(path: str, channel_prefix: str) -> Session:
    """İkili modun ham yedeği (Oskilo çerçeveleri)."""
    from .binary import Decoder
    from .parser import ParsedLine

    dec = Decoder()
    s = Session(mode="file", seq_mode="on", channel_prefix=channel_prefix)
    s.seq32 = True

    def take(frames):
        for fr in frames:
            if fr[0] == "data":
                s.add(float(s.n), ParsedLine("data", values=[float(fr[1])] + fr[2]))
            elif fr[0] != "crc":
                s.add(float(s.n), ParsedLine(fr[0], count=fr[1]))

    with open(path, "rb") as f:
        while True:
            chunk = f.read(1 << 20)
            if not chunk:
                break
            take(dec.feed(chunk))
    take(dec.flush())
    s.crc_errors = dec.crc_errors
    s.flush_pending()
    if s.n == 0:
        raise ValueError("no-data")
    return s


def _split(ln: str, delim: str) -> list[str]:
    """Bir satırı sütunlara böler. Tırnak varsa CSV kuralıyla (virgüllü not gibi), her satır
    kendi başına: elle düzenlenmiş dosyada kapanmamış bir tırnak sonraki satırları yutmaz."""
    if '"' not in ln:
        return ln.split(delim)
    try:
        return next(csv.reader([ln], delimiter=delim, strict=True))
    except (csv.Error, StopIteration):
        return ln.split(delim)


def _load_text(path: str, channel_prefix: str) -> Session:
    body = (ln for ln in _lines(path) if ln.strip() and not ln.lstrip().startswith("#"))
    first = next(body, None)
    if first is not None:
        delim = "\t" if "\t" in first else ("," if "," in first else None)
        if delim:
            heads = [h.strip() for h in _split(first, delim)]
            rows = ([c.strip() for c in _split(ln, delim)] for ln in body)
            s = _session_from_table(heads, rows, channel_prefix)
            if s is not None:
                return s
    # Genel çözümleyici (eski sürüm kayıtları, ham yedekler, düz loglar)
    s = Session(mode="file", seq_mode="auto", channel_prefix=channel_prefix)
    for i, ln in enumerate(_lines(path)):
        p = parse_line(strip_old_timestamp(ln))
        if p.kind in ("empty",):
            continue
        s.add(float(i), p)
    s.flush_pending()
    if s.n == 0:
        raise ValueError("no-data")
    return s
