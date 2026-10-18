"""Oturumu TXT / CSV / Excel (.xlsx) olarak dışa aktarır."""
from __future__ import annotations

import csv
import datetime as _dt
import os
import re
import tempfile
import time

from . import __version__, APP_NAME
from .analysis import convert, label, unit_of
from .i18n import T
from .parser import fmt_num
from .session import Session

FORMATS = {"txt": ".txt", "csv": ".csv", "xlsx": ".xlsx"}
XLSX_MAX_ROWS = 1_000_000          # Excel sınırı 1.048.576; güvenli pay
XLSX_CHART_MAX_ROWS = 20_000
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_BAD_NAME_CH = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
_RESERVED = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}


def safe_text(v: str) -> str:
    """Excel/LibreOffice'in metni formül sanmasını engeller (CSV formül enjeksiyonu).

    Kanal adları karttan gelir; "=HYPERLINK(...)" gibi bir ad dosyayı açanı riske atmasın.
    """
    return "'" + v if v.startswith(_FORMULA_START) else v


def safe_prefix(prefix: str, fallback: str = "deney") -> str:
    """Dosya adı önekini Windows/macOS'ta geçerli ve güvenli hale getirir."""
    p = _BAD_NAME_CH.sub("_", (prefix or "").strip()).replace(" ", "_")
    p = p.strip(". ")[:60]
    if not p or p.upper() in _RESERVED or p in ("..", "."):
        p = fallback
    return p


def _notes_by_row(notes, n: int) -> dict[int, str]:
    out: dict[int, str] = {}
    for row, _t, text in notes:
        r = min(max(0, row), max(0, n - 1))
        out[r] = (out[r] + " | " + text) if r in out else text
    return out


def _columns(s: Session, lang: str, n: int, notes=None, units=None):
    """(başlıklar, satır üreten fonksiyon) döndürür."""
    t = T(lang)
    notes = s.notes if notes is None else notes
    units = s.units if units is None else units
    heads = [t("col_index")]
    getters = [lambda i: i + 1]
    if s.has_seq and len(s.seq) >= n:
        heads.append(t("col_seq"))
        getters.append(lambda i: s.seq[i])
    if s.sample_rate > 0:
        xs, _ = s.x_values(end=n, prefer="time")
        heads.append(t("col_time_s"))
        getters.append(lambda i, xs=xs: round(xs[i], 9))
    if s.mode == "live" and len(s.t_host) >= n and n:
        base = s.started_at
        th = s.t_host
        t0 = th[0]
        heads.append(t("col_pc_time"))
        getters.append(lambda i: (base + _dt.timedelta(seconds=th[i] - t0)))
        heads.append(t("col_elapsed"))
        getters.append(lambda i: round(th[i] - t0, 4))
    if notes:
        by_row = _notes_by_row(notes, n)
        heads.append(t("col_note"))
        getters.append(lambda i, by_row=by_row: by_row.get(i, ""))
    # kanallar her zaman en sonda (Excel grafiği bu sırayı kullanır)
    for name, col in zip(s.channels, s.cols):
        u = unit_of(units, name)
        heads.append(label(name, u))
        if u is None:
            getters.append(lambda i, col=col: col[i])
        else:
            getters.append(lambda i, col=col, u=u: convert(col[i], u))
    return heads, getters


def _fmt_cell(v) -> str:
    if isinstance(v, float):
        return fmt_num(v)
    if isinstance(v, _dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
    if isinstance(v, str):
        return safe_text(v)
    return str(v)


def summary_rows(s: Session, lang: str, n: int | None = None, notes=None, units=None) -> list[tuple[str, str]]:
    t = T(lang)
    n = s.n if n is None else n
    notes = s.notes if notes is None else notes
    units = s.units if units is None else units
    rows = [
        (t("sum_app"), f"{APP_NAME} v{__version__}"),
        (t("sum_mode"), t("mode_" + s.mode)),
    ]
    if s.port:
        rows.append((t("sum_port"), f"{s.port} @ {s.baud} baud"))
    rows += [
        (t("sum_start"), s.started_at.strftime("%Y-%m-%d %H:%M:%S")),
        (t("sum_end"), (s.ended_at or _dt.datetime.now()).strftime("%Y-%m-%d %H:%M:%S")),
        (t("sum_rows"), str(n)),
        (t("sum_channels"), ", ".join(s.channels)),
    ]
    if s.sample_rate > 0:
        rows.append((t("sum_fs"), f"{fmt_num(s.sample_rate)} Hz"))
    if s.expected is not None:
        rows.append((t("sum_expected"), str(s.expected)))
    if s.integrity_known:
        rows.append((t("sum_missing"), str(s.missing)))
        rows.append((t("sum_integrity"), t("ok") if s.integrity_ok else t("not_ok")))
    else:
        rows.append((t("sum_integrity"), t("integrity_unknown")))
    if s.gaps:
        rows.append((t("sum_gaps"), "; ".join(f"{g.after_seq}→{g.next_seq} (-{g.missing})" for g in s.gaps[:50])
                     + (" …" if len(s.gaps) > 50 else "")))
    if s.bad_lines:
        rows.append((t("sum_bad_lines"), str(s.bad_lines)))
    if s.crc_errors:
        rows.append((t("sum_crc"), str(s.crc_errors)))
    if s.fixed_seq:
        rows.append((t("sum_fixed"), str(s.fixed_seq)))
    if s.out_of_order:
        rows.append((t("sum_ooo"), str(s.out_of_order)))
    if s.text_lines:
        rows.append((t("sum_text_lines"), str(s.text_lines)))
    conv = [(name, unit_of(units, name)) for name in s.channels]
    conv = [f"{name}: ×{u['scale']:g} {'+' if u['offset'] >= 0 else '−'} {abs(u['offset']):g}"
            + (f" → {u['unit']}" if u["unit"] else "") for name, u in conv if u]
    if conv:
        rows.append((t("sum_units"), "; ".join(conv)))
    if notes:
        by_row = _notes_by_row(notes, n)
        rows.append((t("sum_notes"), "; ".join(f"#{r + 1}: {txt}" for r, txt in sorted(by_row.items()))[:2000]))
    return rows


def default_filename(s: Session, prefix: str, fmt: str) -> str:
    stamp = s.started_at.strftime("%Y-%m-%d_%H-%M-%S")
    return f"{safe_prefix(prefix)}_{stamp}{FORMATS[fmt]}"


def _replace(tmp: str, path: str) -> None:
    """os.replace; Windows'ta dosya kısa süre kilitliyse (antivirüs, OneDrive) birkaç kez dener."""
    for i in range(10):
        try:
            os.replace(tmp, path)
            return
        except PermissionError as e:
            if i == 9:   # hedef dosya büyük ihtimalle Excel'de açık
                raise PermissionError(e.errno, e.strerror, path) from None
            time.sleep(0.15)


def export_session(s: Session, path: str, fmt: str, lang: str = "tr", info: dict | None = None,
                   units: dict | None = None) -> str:
    """Dosyayı önce geçici adla yazar, bitince asıl ada taşır.

    Böylece yazma yarıda kalırsa (disk doldu, program kapandı) yarım/bozuk bir
    dosya asıl adla ortada kalmaz; varsa eski dosya da bozulmaz.
    """
    path = os.path.abspath(path)
    folder = os.path.dirname(path)
    os.makedirs(folder, exist_ok=True)
    with s.lock:
        if s.ended_at is not None:
            s.flush_pending()
        n = s.n          # kayıt sürerken dışa aktarılırsa: o anki satırlar
        notes = list(s.notes)                 # aynı anın notları (sonradan eklenen bir sonraki kayda kalır)
    units = dict(s.units if units is None else units)
    if info is not None:
        info["rows"] = n
        info["notes"] = len(notes)
    fd, tmp = tempfile.mkstemp(dir=folder, prefix=f".{os.path.basename(path)}.", suffix=".tmp")
    os.close(fd)
    try:
        if fmt == "xlsx":
            _export_xlsx(s, tmp, lang, n, notes, units)
        elif fmt == "csv":
            _export_delim(s, tmp, lang, ",", False, n, notes, units)
        else:
            _export_delim(s, tmp, lang, "\t", True, n, notes, units)
        _replace(tmp, path)
    except BaseException:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise
    return path


def _export_delim(s: Session, path: str, lang: str, delim: str, comments: bool, n: int,
                  notes=None, units=None) -> None:
    heads, getters = _columns(s, lang, n, notes, units)
    with open(path, "w", encoding="utf-8-sig" if not comments else "utf-8", newline="") as f:
        if comments:
            for k, v in summary_rows(s, lang, n, notes, units):
                f.write("# " + f"{k}: {v}".replace("\r", " ").replace("\n", " ") + "\n")
        w = csv.writer(f, delimiter=delim, lineterminator="\n")
        w.writerow([safe_text(str(h)) for h in heads])
        for i in range(n):
            w.writerow([_fmt_cell(g(i)) for g in getters])


def _num(v):
    if isinstance(v, float):
        if v != v:
            return None
        if v.is_integer() and abs(v) < 1e15:
            return int(v)
    return v


def _export_xlsx(s: Session, path: str, lang: str, n: int, notes=None, units=None) -> None:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.chart import Reference, ScatterChart, Series
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    t = T(lang)
    units = s.units if units is None else units
    heads, getters = _columns(s, lang, n, notes, units)
    wb = Workbook(write_only=True)

    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F6FEB")
    center = Alignment(horizontal="center", vertical="center")

    # --- Özet sayfası ----------------------------------------------------
    ws_sum = wb.create_sheet(t("sheet_summary"))
    ws_sum.column_dimensions["A"].width = 26
    ws_sum.column_dimensions["B"].width = 36
    for col in "CDEFG":
        ws_sum.column_dimensions[col].width = 14

    def tcell(ws, v):
        """Metin hücresi: '=' ile başlasa bile formül olarak yazılmaz."""
        c = WriteOnlyCell(ws, value=v)
        if isinstance(v, str) and v.startswith("="):
            c.data_type = "s"
        return c

    def hcell(ws, v):
        c = tcell(ws, v)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = center
        return c

    title = WriteOnlyCell(ws_sum, value=f"{APP_NAME} — {t('sheet_summary')}")
    title.font = Font(bold=True, size=14, color="1F6FEB")
    ws_sum.append([title])
    ws_sum.append([])
    for k, v in summary_rows(s, lang, n, notes, units):
        kc = tcell(ws_sum, k)
        kc.font = Font(bold=True)
        vc = tcell(ws_sum, int(v) if v.isdigit() else v)
        if k == t("sum_integrity"):
            ok = v == t("ok")
            vc.font = Font(bold=True, color="1A7F37" if ok else ("CF222E" if v == t("not_ok") else "9A6700"))
        ws_sum.append([kc, vc])
    ws_sum.append([])
    ws_sum.append([hcell(ws_sum, h) for h in (t("stat_channel"), t("stat_count"), "Min", "Max",
                                              t("stat_mean"), t("stat_std"))])
    for st in s.channel_stats(n):
        u = unit_of(units, st["name"])
        name = label(st["name"], u)
        if st["count"]:
            lo, hi = convert(st["min"], u), convert(st["max"], u)
            if hi < lo:                      # negatif çarpan min/max'ı yer değiştirir
                lo, hi = hi, lo
            std = st["std"] * (abs(u["scale"]) if u else 1)
            ws_sum.append([tcell(ws_sum, name), st["count"], _num(lo), _num(hi),
                           round(convert(st["mean"], u), 6), round(std, 6)])
        else:
            ws_sum.append([tcell(ws_sum, name), 0])

    # --- Veri sayfaları ----------------------------------------------------
    data_sheets = []
    n_parts = max(1, -(-n // XLSX_MAX_ROWS))
    for part in range(n_parts):
        name = t("sheet_data") if n_parts == 1 else f"{t('sheet_data')} {part + 1}"
        ws = wb.create_sheet(name)
        ws.freeze_panes = "A2"
        for ci, h in enumerate(heads, start=1):
            ws.column_dimensions[get_column_letter(ci)].width = max(12, min(28, len(str(h)) + 4))
        ws.append([hcell(ws, h) for h in heads])
        lo, hi = part * XLSX_MAX_ROWS, min(n, (part + 1) * XLSX_MAX_ROWS)
        for i in range(lo, hi):
            row = []
            for g in getters:
                v = g(i)
                row.append(tcell(ws, v) if isinstance(v, str) else _num(v))
            ws.append(row)
        data_sheets.append((ws, hi - lo))

    # --- Grafik ------------------------------------------------------------
    if 0 < n <= XLSX_CHART_MAX_ROWS and s.channels:
        ws_data, rows = data_sheets[0]
        x_col = 1
        for name in (t("col_time_s"), t("col_elapsed"), t("col_seq")):
            if name in heads:
                x_col = heads.index(name) + 1
                break
        first_ch = len(heads) - len(s.channels) + 1
        ws_chart = wb.create_sheet(t("sheet_chart"))
        ch = ScatterChart()
        ch.title = t("chart_title")
        ch.style = 13
        ch.height, ch.width = 12, 26
        ch.x_axis.title = heads[x_col - 1]
        ch.y_axis.title = t("value")
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        xref = Reference(ws_data, min_col=x_col, min_row=2, max_row=rows + 1)
        palette = ["1F6FEB", "2DA44E", "CF222E", "8250DF", "BF8700", "1B7C83", "E16F24", "57606A"]
        for ci in range(first_ch, len(heads) + 1):
            yref = Reference(ws_data, min_col=ci, min_row=1, max_row=rows + 1)
            ser = Series(yref, xref, title_from_data=True)
            ser.marker.symbol = "none"
            ser.smooth = False
            ser.graphicalProperties.line.width = 12700
            ser.graphicalProperties.line.solidFill = palette[(ci - first_ch) % len(palette)]
            ch.series.append(ser)
        ws_chart.add_chart(ch, "A1")

    wb.save(path)
